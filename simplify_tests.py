"""
测试用例精简工具
用法: python3 simplify_tests.py input.xlsx output.xlsx [--header]

参数:
  input.xlsx   输入文件（三列：列1, 用例名称A_B(C), 列3）
  output.xlsx  输出文件
  --header     若原始文件有表头行，加此参数跳过第一行
"""

import sys
import re
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ── 解析 ─────────────────────────────────────────────────────────────────────

def parse_name(name: str):
    """
    从 A_B(C) 格式中提取 A, B, C。
    规则：最后一个 _ 分隔 A 与 B(C)；最后一对括号内为 C。
    返回 (A, B, C)，解析失败返回 (None, None, None)。
    """
    s = str(name).strip()
    last_under = s.rfind("_")
    last_open = s.rfind("(")
    last_close = s.rfind(")")

    if last_under == -1 or last_open == -1 or last_close == -1:
        return None, None, None
    if last_open < last_under or last_close < last_open:
        return None, None, None

    A = s[:last_under]
    B = s[last_under + 1 : last_open]
    C = s[last_open + 1 : last_close]
    return A, B, C


# ── 正交选择 ─────────────────────────────────────────────────────────────────

def select_orthogonal(group_df: pd.DataFrame) -> pd.DataFrame:
    """
    对同一 C 分组：每个 A 保留一行，选择时优先选当前频次最低的 B，
    使 B 的分布尽量均匀（贪心）。
    """
    # A → [row列表]
    a_to_rows: dict[str, list] = defaultdict(list)
    for _, row in group_df.iterrows():
        a_to_rows[row["A"]].append(row)

    # 最受约束的 A（可选 B 种类少的）优先处理
    sorted_a = sorted(
        a_to_rows.items(),
        key=lambda item: len({r["B"] for r in item[1]}),
    )

    b_freq: dict[str, int] = defaultdict(int)
    selected = []

    for _a, rows in sorted_a:
        b_to_rows: dict[str, list] = defaultdict(list)
        for row in rows:
            b_to_rows[row["B"]].append(row)

        # 选频次最低的 B；同频次时按 B 值字母序保证确定性
        best_b = min(b_to_rows.keys(), key=lambda b: (b_freq[b], b))
        selected.append(b_to_rows[best_b][0])
        b_freq[best_b] += 1

    return pd.DataFrame(selected)


# ── 格式化 ───────────────────────────────────────────────────────────────────

# B 值颜色池（最多支持 20 种 B）
_B_COLORS = [
    "FFF2CC", "DDEEFF", "E2EFDA", "FCE4D6", "EDE7F6",
    "E8F5E9", "FFF9C4", "F3E5F5", "E3F2FD", "FBE9E7",
    "D7E4BC", "FFCCCC", "CCE5FF", "CCFFCC", "FFE5CC",
    "E5CCFF", "FFFFE0", "FFD700", "B0E0E6", "F0E68C",
]

_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SUMMARY_FILL = PatternFill("solid", fgColor="D9E1F2")


def _sanitize_sheet_name(name: str) -> str:
    """去除 Excel sheet 名非法字符，限制 31 字符。"""
    return re.sub(r'[\\/:*?"<>|\[\]]', "_", str(name))[:31]


def _format_sheet(ws, b_color_map: dict, is_summary: bool):
    # 表头样式
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行着色（按 B 值）
    if not is_summary:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        b_col_idx = headers.index("B") + 1 if "B" in headers else None
        for row in ws.iter_rows(min_row=2):
            if b_col_idx:
                b_val = row[b_col_idx - 1].value
                color = b_color_map.get(str(b_val) if b_val else "")
                if color:
                    fill = PatternFill("solid", fgColor=color)
                    for cell in row:
                        cell.fill = fill
    else:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.fill = _SUMMARY_FILL

    # 自动列宽
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

    # 冻结首行
    ws.freeze_panes = "A2"


# ── 主流程 ───────────────────────────────────────────────────────────────────

def main(input_path: str, output_path: str, has_header: bool):
    # 读取
    header_row = 0 if has_header else None
    df = pd.read_excel(input_path, header=header_row)
    df.columns = ["col1", "name", "col3"]

    # 解析 A/B/C
    parsed = df["name"].apply(lambda x: pd.Series(parse_name(x), index=["A", "B", "C"]))
    df = pd.concat([df, parsed], axis=1)

    # 过滤无法解析的行
    bad = df[df["C"].isna()]
    if not bad.empty:
        print(f"[警告] 以下 {len(bad)} 行格式无法解析，已跳过：")
        for _, r in bad.iterrows():
            print(f"  {r['name']}")
        df = df.dropna(subset=["C"])

    if df.empty:
        print("[错误] 没有可处理的数据。")
        sys.exit(1)

    # B 值颜色映射（全局，跨 sheet 颜色一致）
    all_b = sorted(df["B"].dropna().unique())
    b_color_map = {b: _B_COLORS[i % len(_B_COLORS)] for i, b in enumerate(all_b)}

    # 写入 Excel
    summary_rows = []
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for c_val, c_group in sorted(df.groupby("C")):
            result = select_orthogonal(c_group)
            result = result.sort_values("A").reset_index(drop=True)

            out = result[["col1", "name", "A", "B", "col3"]].copy()
            out.columns = ["列1", "用例名称", "A", "B", "列3"]

            sheet_name = _sanitize_sheet_name(c_val)
            out.to_excel(writer, sheet_name=sheet_name, index=False)

            summary_rows.append({
                "C（分组）": c_val,
                "原始用例数": len(c_group),
                "精简后用例数": len(result),
                "A 种类数": c_group["A"].nunique(),
                "B 种类数": c_group["B"].nunique(),
                "保留率": f"{len(result)/len(c_group)*100:.0f}%",
            })

        # 汇总 sheet（放最前面）
        summary_df = pd.DataFrame(summary_rows)
        summary_df.to_excel(writer, sheet_name="汇总", index=False)

    # 调整 sheet 顺序：汇总放第一个
    wb = load_workbook(output_path)
    wb.move_sheet("汇总", offset=-len(wb.sheetnames) + 1)

    # 格式化所有 sheet
    for sheet_name in wb.sheetnames:
        _format_sheet(wb[sheet_name], b_color_map, is_summary=(sheet_name == "汇总"))

    wb.save(output_path)

    # 输出统计
    orig_total = sum(r["原始用例数"] for r in summary_rows)
    final_total = sum(r["精简后用例数"] for r in summary_rows)
    print(f"完成！共 {len(summary_rows)} 个分组")
    print(f"原始用例: {orig_total} 条 → 精简后: {final_total} 条（减少 {orig_total-final_total} 条）")
    print(f"输出文件: {output_path}")


DEFAULT_INPUT_PATH = "input.xlsx"
DEFAULT_OUTPUT_PATH = "output.xlsx"

if __name__ == "__main__":
    args = sys.argv[1:]
    has_header = "--header" in args
    paths = [a for a in args if not a.startswith("--")]

    if len(paths) >= 2:
        input_path, output_path = paths[0], paths[1]
    elif len(paths) == 1:
        input_path, output_path = paths[0], DEFAULT_OUTPUT_PATH
    else:
        input_path, output_path = DEFAULT_INPUT_PATH, DEFAULT_OUTPUT_PATH

    main(input_path, output_path, has_header)
